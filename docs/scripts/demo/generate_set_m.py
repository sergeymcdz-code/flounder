#!/usr/bin/env python3
"""
Генератор демонстрационного набора данных «M» для платформы Flounder
(ТЗ демо §7 и §5.4, задание T10 п.1).

Запуск без аргументов:  python3 generate_set_m.py
Выход (каталог out/ рядом со скриптом):
  set-M.xlsx              — 200 заявок, формат шаблона импорта
  set-M-with-errors.xlsx  — те же 200 + 3 намеренно ошибочные строки
  set-M.json              — {athletes: [...], entries: [...]}
  README.md               — описание и сводная статистика

Все имена вымышленные. Генератор детерминирован (seed = 2026).
"""
from __future__ import annotations

import json
import random
import statistics
from collections import Counter
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

SEED = 2026
OUT_DIR = Path(__file__).resolve().parent / "out"

HEADERS = [
    "Фамилия",
    "Имя",
    "Отчество",
    "Дата рождения (ДД.ММ.ГГГГ)",
    "Пол (М/Ж)",
    "Клуб/Город",
    "Дисциплина (код)",
    "Заявочное время (мм:сс.сс или пусто)",
]

N_ENTRIES = 200
N_ATHLETES = 120            # 80 спортсменов на 2 вида + 40 на 1 вид = 200 заявок
N_DOUBLE = N_ENTRIES - N_ATHLETES
YEARS = [2010, 2011, 2012, 2013, 2014, 2015]
NT_SHARE = 0.10
CLUBS = ["СШОР №1", "Дельфин", "Волна", "Нептун", "Акватика", "Старт", "Чайка", "Орбита"]

# Средние результаты на 50 в/с (сек) для мальчиков по году рождения; 100 в/с ≈ ×2.3;
# девочки ≈ на 5 % медленнее. Шум — нормальный.
MEAN_50_BY_YEAR = {2015: 50.0, 2014: 46.8, 2013: 43.6, 2012: 40.4, 2011: 37.2, 2010: 34.0}
SD_50 = 2.2
GIRLS_FACTOR = 1.05
FACTOR_100 = 2.3
# Коридоры (ориентир T10: 50 в/с — 32–50 с, 100 в/с — 1:10–1:55); верхние границы расширены,
# чтобы младшие девочки (2015 г.р., ≈ 2:00 на 100 м) не «прилипали» к границе.
CLIP = {50: (30.0, 56.0), 100: (70.0, 128.0)}

# ---------------------------------------------------------------------------
# Вымышленные имена (без реальных известных людей)
# ---------------------------------------------------------------------------
# Фамилии в мужской форме; женская форма образуется по правилу (см. female_surname).
SURNAMES = [
    "Иванов", "Петров", "Сидоров", "Кузнецов", "Смирнов", "Попов", "Васильев", "Новиков",
    "Морозов", "Волков", "Соловьёв", "Зайцев", "Павлов", "Семёнов", "Голубев", "Виноградов",
    "Богданов", "Воробьёв", "Фёдоров", "Михайлов", "Беляев", "Тарасов", "Белов", "Комаров",
    "Орлов", "Киселёв", "Макаров", "Андреев", "Ковалёв", "Ильин", "Гусев", "Титов",
    "Кузьмин", "Кудрявцев", "Баранов", "Куликов", "Алексеев", "Степанов", "Яковлев", "Сорокин",
    "Сергеев", "Романов", "Захаров", "Борисов", "Королёв", "Герасимов", "Пономарёв", "Григорьев",
    "Лазарев", "Медведев", "Ершов", "Никитин", "Соболев", "Рябов", "Поляков", "Цветков",
    "Данилов", "Жуков", "Фролов", "Журавлёв", "Николаев", "Крылов", "Максимов", "Сидорин",
    "Осипов", "Белоусов", "Федотов", "Дорофеев", "Егоров", "Матвеев", "Бобров", "Дмитриев",
    "Калинин", "Анисимов", "Петухов", "Антонов", "Тимофеев", "Никифоров", "Веселов", "Филиппов",
    "Марков", "Большаков", "Суханов", "Миронов", "Ширяев", "Александров", "Коновалов", "Шестаков",
    "Казаков", "Ефимов", "Денисов", "Громов", "Фомин", "Давыдов", "Мельников", "Щербаков",
    "Блинов", "Колесников", "Карпов", "Афанасьев", "Власов", "Маслов", "Исаков", "Тихонов",
    "Лебединский", "Покровский", "Вишневский", "Загорский", "Островский", "Речной", "Озёрный", "Боровой",
]

MALE_NAMES = [
    "Александр", "Алексей", "Андрей", "Антон", "Арсений", "Артём", "Владимир", "Владислав",
    "Глеб", "Григорий", "Даниил", "Денис", "Дмитрий", "Егор", "Иван", "Илья",
    "Кирилл", "Константин", "Лев", "Максим", "Матвей", "Михаил", "Никита", "Николай",
    "Павел", "Пётр", "Роман", "Семён", "Сергей", "Степан", "Тимофей", "Фёдор", "Ярослав",
]

FEMALE_NAMES = [
    "Александра", "Алиса", "Алина", "Анастасия", "Анна", "Арина", "Валерия", "Варвара",
    "Вера", "Вероника", "Виктория", "Дарья", "Ева", "Екатерина", "Елена", "Елизавета",
    "Злата", "Ирина", "Кира", "Ксения", "Мария", "Марина", "Милана", "Надежда",
    "Наталья", "Ольга", "Полина", "София", "Таисия", "Ульяна", "Юлия", "Яна",
]

# Отчества: (мужская форма, женская форма) — от вымышленного имени отца
PATRONYMICS = [
    ("Александрович", "Александровна"), ("Алексеевич", "Алексеевна"), ("Андреевич", "Андреевна"),
    ("Антонович", "Антоновна"), ("Артёмович", "Артёмовна"), ("Борисович", "Борисовна"),
    ("Вадимович", "Вадимовна"), ("Валерьевич", "Валерьевна"), ("Васильевич", "Васильевна"),
    ("Викторович", "Викторовна"), ("Витальевич", "Витальевна"), ("Владимирович", "Владимировна"),
    ("Владиславович", "Владиславовна"), ("Вячеславович", "Вячеславовна"), ("Геннадьевич", "Геннадьевна"),
    ("Георгиевич", "Георгиевна"), ("Григорьевич", "Григорьевна"), ("Денисович", "Денисовна"),
    ("Дмитриевич", "Дмитриевна"), ("Евгеньевич", "Евгеньевна"), ("Егорович", "Егоровна"),
    ("Иванович", "Ивановна"), ("Игоревич", "Игоревна"), ("Ильич", "Ильинична"),
    ("Кириллович", "Кирилловна"), ("Константинович", "Константиновна"), ("Леонидович", "Леонидовна"),
    ("Максимович", "Максимовна"), ("Михайлович", "Михайловна"), ("Николаевич", "Николаевна"),
    ("Олегович", "Олеговна"), ("Павлович", "Павловна"), ("Петрович", "Петровна"),
    ("Романович", "Романовна"), ("Русланович", "Руслановна"), ("Семёнович", "Семёновна"),
    ("Сергеевич", "Сергеевна"), ("Станиславович", "Станиславовна"), ("Степанович", "Степановна"),
    ("Тимофеевич", "Тимофеевна"), ("Фёдорович", "Фёдоровна"), ("Юрьевич", "Юрьевна"),
    ("Ярославович", "Ярославовна"),
]


def female_surname(m: str) -> str:
    if m.endswith(("ский", "цкий")):
        return m[:-2] + "ая"
    if m.endswith("ой") or m.endswith("ый"):
        return m[:-2] + "ая"
    if m.endswith(("ов", "ев", "ёв", "ин", "ын")):
        return m + "а"
    return m  # несклоняемая


def fmt_time(seconds: float | None) -> str:
    if seconds is None:
        return ""
    cs = int(round(seconds * 100))
    mm, rest = divmod(cs, 6000)
    ss, cc = divmod(rest, 100)
    return f"{mm:02d}:{ss:02d}.{cc:02d}"


def parse_time(s: str) -> float:
    mm, rest = s.split(":")
    return int(mm) * 60 + float(rest)


def random_birth_date(rng: random.Random, year: int) -> date:
    start = date(year, 1, 1)
    ndays = (date(year + 1, 1, 1) - start).days
    return start + timedelta(days=rng.randrange(ndays))


def seed_time(rng: random.Random, year: int, gender: str, dist: int) -> float:
    base = MEAN_50_BY_YEAR[year]
    sd = SD_50
    if gender == "Ж":
        base *= GIRLS_FACTOR
    if dist == 100:
        base *= FACTOR_100
        sd *= FACTOR_100
    t = rng.gauss(base, sd)
    lo, hi = CLIP[dist]
    return round(min(max(t, lo), hi), 2)


def generate() -> tuple[list[dict], list[dict]]:
    rng = random.Random(SEED)

    # --- спортсмены -------------------------------------------------------
    genders = ["М"] * (N_ATHLETES // 2) + ["Ж"] * (N_ATHLETES - N_ATHLETES // 2)
    rng.shuffle(genders)
    years = [YEARS[i % len(YEARS)] for i in range(N_ATHLETES)]  # равномерно по годам
    rng.shuffle(years)

    athletes: list[dict] = []
    used_keys: set[tuple] = set()
    i = 0
    while len(athletes) < N_ATHLETES:
        gender = genders[len(athletes)]
        year = years[len(athletes)]
        surname_m = rng.choice(SURNAMES)
        if gender == "М":
            last, first = surname_m, rng.choice(MALE_NAMES)
            middle = rng.choice(PATRONYMICS)[0]
        else:
            last, first = female_surname(surname_m), rng.choice(FEMALE_NAMES)
            middle = rng.choice(PATRONYMICS)[1]
        bdate = random_birth_date(rng, year)
        key = (last, first, middle, bdate)
        if key in used_keys:
            continue
        used_keys.add(key)
        i += 1
        athletes.append(
            {
                "id": i,
                "lastName": last,
                "firstName": first,
                "middleName": middle,
                "birthDate": bdate.strftime("%d.%m.%Y"),
                "birthYear": year,
                "gender": gender,
                "club": rng.choice(CLUBS),
            }
        )

    # --- заявки -------------------------------------------------------------
    double_ids = set(rng.sample(range(1, N_ATHLETES + 1), N_DOUBLE))
    entries: list[dict] = []
    for a in athletes:
        dists = [50, 100] if a["id"] in double_ids else [rng.choice([50, 100])]
        for d in dists:
            code = f"{'M' if a['gender'] == 'М' else 'F'}-{d}-FREE"
            entries.append({"athleteId": a["id"], "disciplineCode": code, "distanceM": d})

    assert len(entries) == N_ENTRIES
    # 10 % NT, остальные — реалистичное время
    nt_idx = set(rng.sample(range(N_ENTRIES), int(N_ENTRIES * NT_SHARE)))
    by_id = {a["id"]: a for a in athletes}
    for k, e in enumerate(entries):
        a = by_id[e["athleteId"]]
        if k in nt_idx:
            e["seedTime"] = ""
            e["seedTimeMs"] = None
        else:
            t = seed_time(rng, a["birthYear"], a["gender"], e["distanceM"])
            e["seedTime"] = fmt_time(t)
            e["seedTimeMs"] = int(round(t * 1000))

    # перемешать строки, чтобы файл выглядел как живой список заявок
    rng.shuffle(entries)
    for n, e in enumerate(entries, start=1):
        e["row"] = n
    return athletes, entries


def entry_rows(athletes: list[dict], entries: list[dict]) -> list[list[str]]:
    by_id = {a["id"]: a for a in athletes}
    rows = []
    for e in entries:
        a = by_id[e["athleteId"]]
        rows.append(
            [a["lastName"], a["firstName"], a["middleName"], a["birthDate"], a["gender"],
             a["club"], e["disciplineCode"], e["seedTime"]]
        )
    return rows


ERROR_ROWS = [
    # неверная дата
    ["Тестов", "Ошибка", "Датович", "31.02.2013", "М", "Дельфин", "M-50-FREE", "00:40.10"],
    # неизвестный код дисциплины
    ["Тестова", "Ошибка", "Кодовна", "15.06.2012", "Ж", "Волна", "X-999", "00:38.50"],
    # пустая фамилия
    ["", "Безфамильный", "Пустович", "03.03.2014", "М", "Старт", "M-100-FREE", ""],
]


def write_xlsx(path: Path, rows: list[list[str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Заявки"
    ws.append(HEADERS)
    for c in range(1, len(HEADERS) + 1):
        ws.cell(row=1, column=c).font = Font(bold=True)
    for r in rows:
        ws.append(r)
    widths = [18, 14, 18, 26, 10, 14, 18, 36]
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"
    wb.save(path)


# ---------------------------------------------------------------------------
# Проверка результата
# ---------------------------------------------------------------------------
def verify(xlsx_path: Path, athletes: list[dict], entries: list[dict]) -> dict:
    ws = load_workbook(xlsx_path, read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[0])
    assert header == HEADERS, f"Заголовки не совпадают: {header}"
    data = [list(r) for r in rows[1:]]
    assert len(data) == N_ENTRIES, f"Ожидалось {N_ENTRIES} строк, получено {len(data)}"

    keys = {(r[0], r[1], r[2], r[3]) for r in data}
    assert len(keys) == N_ATHLETES, f"Уникальных спортсменов {len(keys)}, ожидалось {N_ATHLETES}"
    assert len(athletes) == N_ATHLETES

    valid_codes = {"M-50-FREE", "F-50-FREE", "M-100-FREE", "F-100-FREE"}
    date_by_person: dict[tuple, str] = {}
    nt = 0
    for r in data:
        last, first, middle, bdate, gender, club, code, t = r
        assert last and first and middle, f"пустое ФИО: {r}"
        assert gender in ("М", "Ж"), r
        assert code in valid_codes, r
        assert (gender == "М") == code.startswith("M-"), f"пол и код не согласованы: {r}"
        assert club in CLUBS, r
        d = date(int(bdate[6:]), int(bdate[3:5]), int(bdate[:2]))  # бросит исключение при неверной дате
        assert d.year in YEARS, r
        assert d.strftime("%d.%m.%Y") == bdate, r
        person = (last, first, middle)
        assert date_by_person.setdefault(person, bdate) == bdate, f"разные даты у {person}"
        if t in (None, ""):
            nt += 1
        else:
            sec = parse_time(t)
            dist = int(code.split("-")[1])
            lo, hi = CLIP[dist]
            assert lo <= sec <= hi, f"время вне диапазона: {r}"
            assert fmt_time(sec) == t, r
    assert nt == int(N_ENTRIES * NT_SHARE), f"NT = {nt}"
    return {"rows": len(data), "unique_athletes": len(keys), "nt": nt}


def verify_errors(xlsx_path: Path) -> None:
    ws = load_workbook(xlsx_path, read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    assert list(rows[0]) == HEADERS
    assert len(rows) - 1 == N_ENTRIES + len(ERROR_ROWS)
    tail = [list(r) for r in rows[-3:]]
    assert tail[0][3] == "31.02.2013" and tail[1][6] == "X-999" and tail[2][0] in (None, "")


# ---------------------------------------------------------------------------
# README со статистикой
# ---------------------------------------------------------------------------
def build_readme(athletes: list[dict], entries: list[dict], check: dict) -> str:
    by_id = {a["id"]: a for a in athletes}
    by_code = Counter(e["disciplineCode"] for e in entries)
    by_year_entries = Counter(by_id[e["athleteId"]]["birthYear"] for e in entries)
    by_year_athletes = Counter(a["birthYear"] for a in athletes)
    by_club_entries = Counter(by_id[e["athleteId"]]["club"] for e in entries)
    by_club_athletes = Counter(a["club"] for a in athletes)
    by_gender = Counter(a["gender"] for a in athletes)
    per_athlete = Counter(e["athleteId"] for e in entries)
    n_double = sum(1 for v in per_athlete.values() if v == 2)
    nt = sum(1 for e in entries if not e["seedTime"])

    def time_stats(code: str) -> str:
        ts = [e["seedTimeMs"] / 1000 for e in entries if e["disciplineCode"] == code and e["seedTimeMs"]]
        return (f"{fmt_time(min(ts))} / {fmt_time(statistics.mean(ts))} / {fmt_time(max(ts))} "
                f"(n = {len(ts)})")

    def time_by_year(dist: int) -> list[str]:
        out = []
        for y in YEARS:
            for g in ("М", "Ж"):
                ts = [e["seedTimeMs"] / 1000 for e in entries
                      if e["distanceM"] == dist and e["seedTimeMs"]
                      and by_id[e["athleteId"]]["birthYear"] == y
                      and by_id[e["athleteId"]]["gender"] == g]
                if ts:
                    out.append(f"| {y} | {g} | {len(ts)} | {fmt_time(statistics.mean(ts))} | "
                               f"{fmt_time(min(ts))} | {fmt_time(max(ts))} |")
        return out

    L = []
    L.append("# Демо-набор данных «M» (Flounder, T10 п.1)\n")
    L.append("Сгенерировано скриптом `scripts/demo/generate_set_m.py` (Python 3, openpyxl), "
             f"seed = {SEED}. Все ФИО вымышленные, совпадения случайны.\n")
    L.append("## Файлы\n")
    L.append("| Файл | Содержимое |")
    L.append("|---|---|")
    L.append("| `set-M.xlsx` | 200 заявок в формате шаблона импорта (ТЗ §5.4), лист «Заявки» |")
    L.append("| `set-M-with-errors.xlsx` | те же 200 строк + 3 ошибочные строки в конце (строки 202–204 файла): "
             "неверная дата `31.02.2013`, неизвестный код `X-999`, пустая фамилия |")
    L.append("| `set-M.json` | `athletes[]` (120 спортсменов) и `entries[]` (200 заявок, ссылка `athleteId`) |")
    L.append("| `README.md` | этот файл |\n")
    L.append("## Формат Excel\n")
    L.append("Первая строка — заголовки, строго по ТЗ §5.4:\n")
    L.append("```")
    L.append(" | ".join(HEADERS))
    L.append("```\n")
    L.append("Одна строка = одна заявка. Дата — `ДД.ММ.ГГГГ`, пол — `М`/`Ж`, время — `мм:сс.сс` "
             "или пусто (NT). Коды дисциплин: `M-50-FREE`, `F-50-FREE`, `M-100-FREE`, `F-100-FREE` "
             "(мальчики только `M-`, девочки только `F-`).\n")
    L.append("## Параметры генерации\n")
    L.append(f"- Спортсменов: {len(athletes)} (М: {by_gender['М']}, Ж: {by_gender['Ж']}); "
             f"из них {n_double} заявлены на оба вида (50 и 100 в/с), {len(athletes) - n_double} — на один.")
    L.append("- Годы рождения 2010–2015, равномерно (20 спортсменов на год). "
             "Дата рождения — случайный день года; у одного спортсмена дата одинакова во всех строках.")
    L.append(f"- Клубы (8): {', '.join(CLUBS)}; города не используются.")
    L.append("- Заявочное время: нормальное распределение вокруг средних по году рождения "
             "(50 в/с: 2015 г.р. ≈ 50 с … 2010 г.р. ≈ 34 с, σ ≈ 2.2 с; 100 в/с ≈ ×2.3; девочки ≈ +5 %), "
             f"обрезка 50 в/с: {CLIP[50][0]:.0f}–{CLIP[50][1]:.1f} с, 100 в/с: {fmt_time(CLIP[100][0])}–{fmt_time(CLIP[100][1])}.")
    L.append(f"- Без времени (NT): {nt} заявок ({nt / len(entries):.0%}).\n")
    L.append("## Сводная статистика\n")
    L.append("### По кодам дисциплин\n")
    L.append("| Код | Заявок | NT | Время мин / среднее / макс |")
    L.append("|---|---|---|---|")
    for code in ["M-50-FREE", "F-50-FREE", "M-100-FREE", "F-100-FREE"]:
        nt_c = sum(1 for e in entries if e["disciplineCode"] == code and not e["seedTime"])
        L.append(f"| {code} | {by_code[code]} | {nt_c} | {time_stats(code)} |")
    L.append("")
    L.append("### По годам рождения\n")
    L.append("| Год | Спортсменов | Заявок |")
    L.append("|---|---|---|")
    for y in YEARS:
        L.append(f"| {y} | {by_year_athletes[y]} | {by_year_entries[y]} |")
    L.append("")
    L.append("### Заявочные времена по году и полу (среднее / мин / макс)\n")
    for dist in (50, 100):
        L.append(f"**{dist} в/с**\n")
        L.append("| Год | Пол | n | Среднее | Мин | Макс |")
        L.append("|---|---|---|---|---|---|")
        L.extend(time_by_year(dist))
        L.append("")
    L.append("### По клубам\n")
    L.append("| Клуб | Спортсменов | Заявок |")
    L.append("|---|---|---|")
    for club in CLUBS:
        L.append(f"| {club} | {by_club_athletes[club]} | {by_club_entries[club]} |")
    L.append("")
    L.append("### Доля NT\n")
    L.append(f"{nt} из {len(entries)} заявок без заявочного времени — {nt / len(entries):.1%}.\n")
    L.append("## Проверка\n")
    L.append("Скрипт после генерации перечитывает `set-M.xlsx` и проверяет: заголовки символ в символ, "
             f"{check['rows']} строк, {check['unique_athletes']} уникальных спортсменов (ФИО + дата), "
             "согласованность пола и кода дисциплины, корректность дат, разбор всех времён и их "
             f"попадание в диапазон, ровно {check['nt']} NT. В `set-M-with-errors.xlsx` проверяется наличие "
             "трёх ошибочных строк в конце.\n")
    L.append("## Ожидаемый результат импорта\n")
    L.append("- `set-M.xlsx`: `total: 200, created: 200, errors: []`, `createdProfiles: 120`.")
    L.append("- `set-M-with-errors.xlsx`: `total: 203`, 3 ошибки — строки 202 (Дата рождения), "
             "203 (Дисциплина (код)), 204 (Фамилия); остальные 200 импортируются.")
    return "\n".join(L) + "\n"


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    athletes, entries = generate()
    rows = entry_rows(athletes, entries)

    xlsx = OUT_DIR / "set-M.xlsx"
    xlsx_err = OUT_DIR / "set-M-with-errors.xlsx"
    write_xlsx(xlsx, rows)
    write_xlsx(xlsx_err, rows + ERROR_ROWS)

    with (OUT_DIR / "set-M.json").open("w", encoding="utf-8") as f:
        json.dump(
            {
                "meta": {"seed": SEED, "generator": "scripts/demo/generate_set_m.py",
                         "clubs": CLUBS, "headers": HEADERS},
                "athletes": athletes,
                "entries": entries,
            },
            f, ensure_ascii=False, indent=2,
        )

    check = verify(xlsx, athletes, entries)
    verify_errors(xlsx_err)
    (OUT_DIR / "README.md").write_text(build_readme(athletes, entries, check), encoding="utf-8")

    print(f"OK: {check['rows']} заявок, {check['unique_athletes']} спортсменов, NT = {check['nt']}")
    print(" | ".join(HEADERS))
    for r in rows[:5]:
        print(" | ".join(r))


if __name__ == "__main__":
    main()
